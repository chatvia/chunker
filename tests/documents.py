"""Builders for small, real documents used as test fixtures."""

import io

import docx
import pptx


def _pdf_escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf(pages: list[str]) -> bytes:
    """A minimal text PDF with one page per entry (lines separated by newlines)."""
    page_ids = [4 + 2 * i for i in range(len(pages))]
    kids = " ".join(f"{page_id} 0 R" for page_id in page_ids)
    objects: dict[int, bytes] = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        2: f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode(),
        3: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    }
    for page_id, text in zip(page_ids, pages, strict=True):
        content_id = page_id + 1
        objects[page_id] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_id} 0 R >>"
        ).encode()
        operations = ["BT", "/F1 12 Tf", "72 720 Td", "16 TL"]
        operations += [f"({_pdf_escape(line)}) Tj T*" for line in text.split("\n")]
        operations.append("ET")
        stream = "\n".join(operations).encode()
        objects[content_id] = b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream)

    output = bytearray(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}
    for object_id in sorted(objects):
        offsets[object_id] = len(output)
        output += b"%d 0 obj\n%s\nendobj\n" % (object_id, objects[object_id])
    xref_offset = len(output)
    size = max(objects) + 1
    output += b"xref\n0 %d\n0000000000 65535 f \n" % size
    for object_id in range(1, size):
        output += b"%010d 00000 n \n" % offsets[object_id]
    output += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        size,
        xref_offset,
    )
    return bytes(output)


def make_docx(sections: list[tuple[str, str]]) -> bytes:
    """A Word document with a heading and a paragraph per (heading, body) pair."""
    document = docx.Document()
    for heading, body in sections:
        document.add_heading(heading, level=1)
        document.add_paragraph(body)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_pptx(slides: list[tuple[str, str]]) -> bytes:
    """A PowerPoint deck with a title and body text per (title, body) pair."""
    presentation = pptx.Presentation()
    layout = presentation.slide_layouts[1]  # "Title and Content"
    for title, body in slides:
        slide = presentation.slides.add_slide(layout)
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()
